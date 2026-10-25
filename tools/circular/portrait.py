import re, sys, copy
from pptx import Presentation
from pptx.util import Emu
from lxml import etree

SRC, OUT = sys.argv[1], sys.argv[2]
import os
BODY = float(os.environ.get('BODYPT','10.5'))
A = 'http://schemas.openxmlformats.org/drawingml/2006/main'
q = lambda t: '{%s}%s' % (A, t)

p = Presentation(SRC)
W, H = 7560000, 10692000
p.slide_width, p.slide_height = Emu(W), Emu(H)
s = p.slides[0]
S = {sh.shape_id: sh for sh in s.shapes}
M = 420000            # one margin for every left and right edge
L, R = M, W - M
CW = R - L

def place(sid, x, y, w=None, h=None):
    sh = S[sid]
    sh.left, sh.top = Emu(int(x)), Emu(int(y))
    if w is not None: sh.width = Emu(int(w))
    if h is not None: sh.height = Emu(int(h))

def scale_sz(sid, mp):
    for e in S[sid]._element.iter():
        v = e.get('sz')
        if v and e.tag.endswith(('}rPr', '}defRPr', '}endParaRPr')):
            v = int(v); e.set('sz', str(mp.get(v, v)))

def zero_inset(sid):
    bp = S[sid]._element.find('.//' + q('bodyPr'))
    for k in ('lIns', 'rIns', 'tIns', 'bIns'): bp.set(k, '0')

def ppr_of(pp):
    ppr = pp.find(q('pPr'))
    if ppr is None:
        ppr = etree.Element(q('pPr')); pp.insert(0, ppr)
    return ppr

def line_pts(sid, pts):
    for i, pp in enumerate(S[sid]._element.iter(q('p'))):
        if i >= len(pts): break
        ppr = ppr_of(pp)
        for old in ppr.findall(q('lnSpc')): ppr.remove(old)
        ln = etree.Element(q('lnSpc')); etree.SubElement(ln, q('spcPts')).set('val', str(pts[i]))
        ppr.insert(0, ln)

def align(sid, a):
    for pp in S[sid]._element.iter(q('p')): ppr_of(pp).set('algn', a)

# ---- script cell back to the template format: name (code) ticker ----
runs = list(S[5]._element.iter(q('r')))
for i, r in enumerate(runs):
    t = r.find(q('t'))
    if t is not None and t.text == 'عقار - AQAR`R':
        vals = ['عقار (', '421', ') ', 'AQAR']
        for j, v in enumerate(vals):
            tt = runs[i + j].find(q('t')); tt.text = v
            if v != v.strip(): tt.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
        break
else:
    pass  # template mode: script text already filled

# ================= header =================
HH = 1320000
place(9, 0, 0, W, HH)
k = 0.70
l22, t22, w22, h22 = S[22].left, S[22].top, S[22].width, S[22].height
l26, t26, w26, h26 = S[26].left, S[26].top, S[26].width, S[26].height
mx = R - w26 * k; my = (HH - h26 * k) / 2
place(26, mx, my, w26 * k, h26 * k)
place(22, mx - (l26 - l22) * k, my - (t26 - t22) * k, w22 * k, h22 * k)
gap = 220000
TW = (mx - gap - L - gap) / 2
TY, TH = 190000, HH - 300000
place(10, L, TY, TW, TH)
place(8, L + TW + gap, TY, TW, TH)
for sid in (8, 10):
    scale_sz(sid, {3200: 2400, 1800: 1150, 1600: 1150})
    zero_inset(sid)
    line_pts(sid, [3000, 1650, 1650, 1650])
place(17, 0, HH, W * 0.78, 70000)
place(20, W * 0.78, HH, W * 0.22, 70000)

# ================= date line =================
y = HH + 200000
zero_inset(31); zero_inset(33)
place(31, L + 90000, y)
S[33]._element.find('.//'+q('bodyPr')).set('wrap','square')
place(33, R - 90000 - 1600000, y, 1600000)
align(33, 'r')
place(27, L, y + 230000, CW, 0)

# ================= table =================
ty = y + 330000
tbl = S[5].table
cw = [c.width for c in tbl.columns]; tot = sum(cw)
ws = [int(w0 * CW / tot) for w0 in cw]; ws[-1] += CW - sum(ws)
for c, w0 in zip(tbl.columns, ws): c.width = Emu(w0)
rh = [340000, 280000, 240000, 240000, 420000, 260000, 560000, 280000, 280000, 560000]
for r, h in zip(tbl.rows, rh): r.height = Emu(h)
place(5, L, ty, CW, sum(rh))
scale_sz(5, {1600: 1500, 1400: 1250, 1200: 1100})

# ================= footer =================
FH = 760000; fy = H - FH
place(35, 0, fy, W, FH)
place(37, 0, fy, W * 0.3, 50000); place(39, W * 0.3, fy, W * 0.7, 50000)
g, t41 = S[53], S[42]
dx, dy = t41.left - g.left, t41.top - g.top          # original flag/towers ↔ company-name offset
gx, gy = L, fy + 60000
place(53, gx, gy)
place(42, gx + dx, gy + dy, 3300000, t41.height)
qs = 560000
qx, qy = R - qs, fy + 75000
place(44, qx, qy, qs, qs)
cap_w = 1000000
place(52, qx + qs / 2 - cap_w / 2, qy + qs + 15000, cap_w, 160000)
zero_inset(52); align(52, 'ctr')

# ================= letter body =================
by = ty + sum(rh) + 250000
fb = fy - 190000
pad = 170000
colgap = 360000
inner_l = L + pad
cwid = (R - inner_l - colgap) / 2
xL, xR = inner_l, inner_l + cwid + colgap
for sid in (16, 2, 6, 3): zero_inset(sid)
place(16, xL, by, cwid, 300000)
place(2, xR, by, cwid, 300000)
place(6, xL, by + 330000, cwid, fb - by - 400000)
place(3, xR, by + 330000, cwid, fb - by - 400000)
place(24, inner_l + cwid + colgap / 2, by + 60000, 0, fb - by - 160000)
place(54, L, by + 20000, 0, fb - by - 20000)
place(56, L, fb, CW, 0)
b = int(BODY * 100)
mp = {1300: b, 1400: b + 100, 1600: b + 250, 2400: b + 500, 800: 700, 1100: 900, 700: 600, 200: 200}
for sid in (3, 6): scale_sz(sid, mp)
for sid in (2, 16): scale_sz(sid, {1600: b + 250})

# ================= one numeral style everywhere =================
# GE SS draws 0-9 as Arabic-Indic shapes and has no Latin letters, '@' or '%'.
# Digits / Latin / @ / % inside GE SS runs are split out and set in Eurostile.
pat = re.compile(r'([0-9A-Za-z@%][0-9A-Za-z@%/.:,\-– +]*[0-9A-Za-z@%]|[0-9A-Za-z@%])')
XS = '{http://www.w3.org/XML/1998/namespace}space'
for sh in s.shapes:
    for r in list(sh._element.iter(q('r'))):
        t = r.find(q('t'))
        if t is None or not t.text or not pat.search(t.text): continue
        rpr = r.find(q('rPr'))
        cs = rpr.find(q('cs')) if rpr is not None else None
        if cs is None or not cs.get('typeface', '').startswith('GE SS'): continue
        parts = [x for x in pat.split(t.text) if x]
        parent = r.getparent(); idx = parent.index(r); parent.remove(r)
        for i, part in enumerate(parts):
            nr = copy.deepcopy(r); nt = nr.find(q('t')); nt.text = part
            if pat.fullmatch(part):
                c2 = nr.find(q('rPr')).find(q('cs')); c2.set('typeface', 'Eurostile'); c2.attrib.pop('panose', None)
            if part != part.strip(): nt.set(XS, 'preserve')
            parent.insert(idx + i, nr)

p.save(OUT)
print('saved', OUT)

# ================= paired bilingual body: one borderless table, EN | AR per paragraph pair =================
import os
p3 = Presentation(OUT); s3 = p3.slides[0]
S3 = {sh.shape_id: sh for sh in s3.shapes}
def paras(sid): return list(S3[sid].text_frame._txBody.iter(q('p')))
EN, AR = paras(6), paras(3)
SAL_EN, SAL_AR = paras(16), paras(2)
groups = [(SAL_EN, SAL_AR), ([EN[0]], [AR[0]]), ([EN[2]], [AR[2]]), ([EN[4]], [AR[4]]), ([EN[7]], [AR[6]]),
          (EN[9:13], AR[8:12]), (EN[14:18], AR[13:17])]
BY = int(os.environ.get('BY', '0')) or S3[16].top
TL = L
gapc = 300000
colw = (CW - gapc) // 2
shape = s3.shapes.add_table(len(groups), 2, Emu(TL), Emu(BY), Emu(CW), Emu(len(groups) * 200000))
shape.name = 'Body pairs'
tb = shape.table
tb.columns[0].width = Emu(colw + gapc // 2); tb.columns[1].width = Emu(CW - colw - gapc // 2)
tblPr = shape._element.find('.//' + q('tblPr'))
for k_ in list(tblPr.attrib): del tblPr.attrib[k_]
for ch in list(tblPr): tblPr.remove(ch)
def noline(tag):
    e = etree.Element(q(tag)); e.set('w', '0'); etree.SubElement(e, q('noFill')); return e
for ri, (en_ps, ar_ps) in enumerate(groups):
    tb.rows[ri].height = Emu(150000)
    for ci, ps in enumerate((en_ps, ar_ps)):
        cell = tb.cell(ri, ci)
        txb = cell._tc.find(q('txBody'))
        for old in txb.findall(q('p')): txb.remove(old)
        for pp in ps:
            np_ = copy.deepcopy(pp)
            ppr = ppr_of(np_)
            ppr.set('algn', 'l' if ci == 0 else 'r')
            for t in np_.iter(q('t')):
                if t.text and ci == 1: t.text = t.text.replace('ـ', '')   # drop manual tatweel stretching
            txb.append(np_)
        tcPr = cell._tc.find(q('tcPr'))
        if tcPr is None: tcPr = etree.SubElement(cell._tc, q('tcPr'))
        for ch in list(tcPr): tcPr.remove(ch)
        last = ri == len(groups) - 1
        tcPr.set('marL', str(170000 if ci == 0 else gapc // 2))
        tcPr.set('marR', str(gapc // 2 if ci == 0 else 0))
        tcPr.set('marT', '0'); tcPr.set('marB', str(0 if last else (70000 if ri == 0 else 105000)))
        for tag in ('lnL', 'lnR', 'lnT', 'lnB'):
            if tag == 'lnR' and ci == 0 and ri > 0:
                e = etree.Element(q('lnR')); e.set('w', '6350')
                sf = etree.SubElement(e, q('solidFill')); etree.SubElement(sf, q('srgbClr')).set('val', 'D9DEE1')
                tcPr.append(e)
            else:
                tcPr.append(noline(tag))
        etree.SubElement(tcPr, q('noFill'))
# remove the old separate text boxes and their divider
for sid in (6, 3, 16, 2, 24):
    el = S3[sid]._element; el.getparent().remove(el)
# bracket: vertical from body top to bracket bottom; bottom line placed from env (measured after render)
FB = int(os.environ.get('FB', '0'))
if FB:
    S3[56].top = Emu(FB); S3[54].height = Emu(FB - S3[54].top)
p3.save(OUT)
print('body table done')
